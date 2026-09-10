#!/usr/bin/env python3
"""Phase 5 loader, 2026-09-10: fetch the odd-shaped sources to a local csv, then
hand the file to cms_years_fast_load.land() for gzip parts, PUT, one COPY INTO.

Two fetch kinds, both new:
  zip_member_ranged  read one member out of a remote zip with Range requests:
                     central directory from the tail, local header at the member
                     offset, inflate that member only. EOIR's zip is 4.56 GB and
                     we want two members.
  dol_api_paged      page apiprod.dol.gov v4 at 10k rows a call into one csv,
                     sorted by the key so pages do not overlap. Backs off on 429.

    python scripts/phase5_load.py --sid FED_EOIR_JUDGE            # fetch + count only
    python scripts/phase5_load.py --sid FED_EOIR_JUDGE --run      # land
"""
from __future__ import annotations

import argparse
import csv
import io
import os
import re
import struct
import sys
import time
import zlib
from pathlib import Path

import requests

_REPO = Path(__file__).resolve().parents[1]
for p in (_REPO, _REPO / "scripts", _REPO / "library-onboarding"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))
try:
    from dotenv import load_dotenv
    load_dotenv(_REPO / "library-onboarding" / ".env", override=True)
except Exception:  # pragma: no cover
    pass

import snow  # noqa: E402
import bridge_fuel_load as bf  # noqa: E402
import cms_years_fast_load as fast  # noqa: E402
import sprint_phase5_specs as SPECMOD  # noqa: E402
import sprint_phase2_specs as SPECMOD2  # noqa: E402

UA = {"User-Agent": "Ripple-Library/1.0 (data onboarding; w.rogers9999@gmail.com)"}
SCRATCH = _REPO / ".scratch" / "phase5"
SCRATCH.mkdir(parents=True, exist_ok=True)
log = fast.log
csv.field_size_limit(1 << 30)


# ------------------------------------------------------------ ranged zip member
def _central_directory(url: str, size: int, tail_bytes: int = 2_000_000) -> dict[str, tuple[int, int, int]]:
    """name -> (compressed size, uncompressed size, local header offset). Zip64 offsets honoured."""
    tail = requests.get(url, headers={**UA, "Range": f"bytes={size - tail_bytes}-{size - 1}"}, timeout=180).content
    out = {}
    for m in re.finditer(rb"PK\x01\x02", tail):
        o = m.start()
        h = tail[o:o + 46]
        if len(h) < 46:
            continue
        comp, uncomp = struct.unpack("<II", h[20:28])
        nlen, xlen = struct.unpack("<HH", h[28:32])
        off = struct.unpack("<I", h[42:46])[0]
        name = tail[o + 46:o + 46 + nlen].decode("latin-1")
        extra = tail[o + 46 + nlen:o + 46 + nlen + xlen]
        # zip64 extra: fields present only for the values that overflowed, in order uncomp, comp, offset
        i = 0
        while i + 4 <= len(extra):
            hid, hlen = struct.unpack("<HH", extra[i:i + 4])
            body = extra[i + 4:i + 4 + hlen]
            if hid == 1:
                j = 0
                if uncomp == 0xFFFFFFFF:
                    uncomp = struct.unpack("<Q", body[j:j + 8])[0]
                    j += 8
                if comp == 0xFFFFFFFF:
                    comp = struct.unpack("<Q", body[j:j + 8])[0]
                    j += 8
                if off == 0xFFFFFFFF:
                    off = struct.unpack("<Q", body[j:j + 8])[0]
                    j += 8
            i += 4 + hlen
        out[name] = (comp, uncomp, off)
    return out


def fetch_zip_member(url: str, member_suffix: str, dest: Path) -> Path:
    if dest.exists() and dest.stat().st_size > 0:
        log(f"    member cached: {dest.name} ({dest.stat().st_size/1e9:.2f} GB)")
        return dest
    size = int(requests.head(url, headers=UA, timeout=120, allow_redirects=True).headers["Content-Length"])
    cd = _central_directory(url, size)
    hits = [n for n in cd if n.endswith(member_suffix)]
    if len(hits) != 1:
        raise RuntimeError(f"member {member_suffix!r}: {len(hits)} matches in {len(cd)} entries")
    name = hits[0]
    comp, uncomp, off = cd[name]
    lh = requests.get(url, headers={**UA, "Range": f"bytes={off}-{off + 29}"}, timeout=120).content
    if lh[:4] != b"PK\x03\x04":
        raise RuntimeError(f"bad local header at {off} for {name}")
    method = struct.unpack("<H", lh[8:10])[0]
    lnlen, lxlen = struct.unpack("<HH", lh[26:30])
    start = off + 30 + lnlen + lxlen
    log(f"    member {name}: {comp/1e6:.1f} MB compressed, {uncomp/1e9:.2f} GB raw, method {method}")
    d = zlib.decompressobj(-15) if method == 8 else None
    tmp = dest.with_suffix(".part")
    got = 0
    with requests.get(url, headers={**UA, "Range": f"bytes={start}-{start + comp - 1}"},
                      stream=True, timeout=(30, 300)) as r, open(tmp, "wb") as f:
        if r.status_code != 206:
            raise RuntimeError(f"Range not honoured: {r.status_code}")
        for block in r.iter_content(1 << 20):
            if not block:
                continue
            got += len(block)
            f.write(d.decompress(block) if d else block)
        if d:
            f.write(d.flush())
    if got != comp:
        raise RuntimeError(f"short member read: {got} of {comp}")
    os.replace(tmp, dest)
    log(f"    member written: {dest.stat().st_size/1e9:.2f} GB")
    return dest


# ------------------------------------------------------------ DOL v4 paged API
def fetch_dol_paged(spec: dict, dest: Path) -> Path:
    if dest.exists() and dest.stat().st_size > 0:
        log(f"    api csv cached: {dest.name}")
        return dest
    key = os.environ.get("DOL_API_KEY", "").strip()
    if not key:
        raise RuntimeError("DOL_API_KEY missing from library-onboarding/.env")
    base = spec["download_url"]
    page = int(spec.get("page_rows", 10_000))
    sort_by = spec.get("sort_by", "")
    tmp = dest.with_suffix(".part")
    header = None
    total = 0
    offset = 0
    wait = 3.0
    with open(tmp, "w", encoding="utf-8", newline="") as out:
        w = csv.writer(out, lineterminator="\n")
        while True:
            q = f"{base}?limit={page}&offset={offset}"
            if sort_by:
                q += f"&sort_by={sort_by}&sort=asc"
            q += f"&X-API-KEY={key}"
            r = None
            for attempt in range(1, 9):
                r = requests.get(q, headers=UA, timeout=300)
                if r.status_code == 429:
                    back = min(600, 30 * attempt)
                    log(f"    429 at offset {offset:,}, sleeping {back}s")
                    time.sleep(back)
                    continue
                break
            if r.status_code == 500 and offset > 0:
                break  # past the end: DOL answers 500, not an empty page (probed 2026-09-10)
            r.raise_for_status()
            rows = list(csv.reader(io.StringIO(r.text, newline="")))
            if not rows or len(rows) < 2:
                break
            if header is None:
                header = rows[0]
                w.writerow(header)
            elif rows[0] != header:
                raise RuntimeError(f"header drift at offset {offset}")
            for rec in rows[1:]:
                w.writerow(rec)
            n = len(rows) - 1
            total += n
            log(f"    page offset {offset:,}: {n:,} rows, total {total:,}")
            offset += page
            if n < page:
                break
            time.sleep(wait)
    os.replace(tmp, dest)
    return dest


# ------------------------------------------------------------ repair
def repair_rows(src: Path, spec: dict) -> Path:
    """EOIR tab files carry unquoted tabs and newlines inside text fields.
    A newline splits one row into two short fragments; a tab makes a row one
    field too wide. Rejoin consecutive short fragments whose widths sum back to
    the header width (the newline becomes a space). Rows still the wrong width
    go to <sid>.rejects.tsv and are counted in the load message, never landed."""
    delim = spec.get("delimiter", ",")
    enc = spec.get("encoding", "utf-8-sig")
    out = src.with_name(src.stem + ".clean.csv")
    rej = src.with_name(src.stem + ".rejects.tsv")
    if out.exists() and out.stat().st_size > 0:
        log(f"    repaired file cached: {out.name}")
        return out
    n_in = n_out = n_join = n_rej = 0
    with (open(src, encoding=enc, errors="replace", newline="") as f,
          open(out, "w", encoding="utf-8", newline="") as fo,
          open(rej, "w", encoding="utf-8", newline="") as fr):
        rdr = csv.reader(f, delimiter=delim, quoting=csv.QUOTE_NONE)
        w = csv.writer(fo, delimiter=delim, lineterminator="\n", quoting=csv.QUOTE_NONE, quotechar=None)
        wr = csv.writer(fr, delimiter=delim, lineterminator="\n", quoting=csv.QUOTE_NONE, quotechar=None)
        header = next(rdr)
        width = len(header)
        w.writerow(header)
        pending = None
        for rec in rdr:
            if not rec:
                continue
            n_in += 1
            if pending is not None:
                joined = pending[:-1] + [pending[-1] + " " + rec[0]] + rec[1:]
                if len(joined) == width:
                    w.writerow(joined); n_out += 1; n_join += 1; pending = None
                    continue
                if len(joined) < width:
                    pending = joined
                    continue
                wr.writerow(pending); n_rej += 1; pending = None
            if len(rec) == width:
                w.writerow(rec); n_out += 1
            elif len(rec) < width:
                pending = rec
            else:
                wr.writerow(rec); n_rej += 1
        if pending is not None:
            wr.writerow(pending); n_rej += 1
    log(f"    repair: {n_in:,} rows read, {n_out:,} kept, {n_join:,} rejoined from newline splits, "
        f"{n_rej:,} rejected to {rej.name}")
    spec["_repair_note"] = f"{n_join} rows rejoined across embedded newlines, {n_rej} rows quarantined for stray tabs"
    return out


# ------------------------------------------------------------ xlsx, bare or in a zip
def _xlsx_to_csv(wb, spec: dict, dest: Path, tag: str) -> None:
    """Write the chosen sheets of an openpyxl workbook as one comma csv.
    skip_rows drops title rows above the header. sheets: None = active sheet,
    "all" = every sheet, unioned on the superset of headers, with a SHEET_NAME
    column first. Every cell as text, dates as YYYY-MM-DD."""
    skip = int(spec.get("skip_rows", 0))
    want = spec.get("sheets")
    if want == "all":
        names = list(wb.sheetnames)
    elif want:
        names = list(want)
    elif spec.get("sheet"):
        names = [spec["sheet"]]
    else:
        names = [wb.active.title]

    def cell(v):
        if v is None:
            return ""
        if hasattr(v, "strftime"):
            return v.strftime("%Y-%m-%d")
        return str(v)

    tables = []
    for nm in names:
        ws = wb[nm]
        it = ws.iter_rows(values_only=True)
        for _ in range(skip):
            next(it, None)
        header = [cell(h).strip() for h in next(it)]
        while header and header[-1] == "":
            header.pop()
        rows = []
        for row in it:
            vals = [cell(v) for v in row[:len(header)]]
            if any(vals):
                rows.append(vals)
        tables.append((nm, header, rows))
    union = []
    for _, h, _ in tables:
        for c in h:
            if c and c not in union:
                union.append(c)
    n = 0
    with open(dest, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        multi = len(tables) > 1
        w.writerow((["SHEET_NAME"] if multi else []) + union)
        for nm, h, rows in tables:
            pos = {c: i for i, c in enumerate(h)}
            for r in rows:
                out = [r[pos[c]] if c in pos and pos[c] < len(r) else "" for c in union]
                w.writerow(([nm] if multi else []) + out)
                n += 1
    log(f"    xlsx {tag}: sheets {names}, {n:,} data rows, {len(union)} columns")


def fetch_zip_xlsx(spec: dict, dest: Path) -> Path:
    import zipfile
    import openpyxl
    if dest.exists() and dest.stat().st_size > 0:
        log(f"    xlsx csv cached: {dest.name}")
        return dest
    zpath = dest.with_suffix(".zip")
    fast.download(spec["download_url"], zpath)
    with zipfile.ZipFile(zpath) as z:
        pat = spec.get("member")
        members = [n for n in z.namelist() if n.lower().endswith(".xlsx") and (not pat or re.search(pat, n))]
        if len(members) != 1:
            raise RuntimeError(f"expected one xlsx member for {pat!r}, found {members}")
        wb = openpyxl.load_workbook(io.BytesIO(z.read(members[0])), read_only=True)
    _xlsx_to_csv(wb, spec, dest, members[0])
    return dest


def fetch_url_xlsx(spec: dict, dest: Path) -> Path:
    import openpyxl
    if dest.exists() and dest.stat().st_size > 0:
        log(f"    xlsx csv cached: {dest.name}")
        return dest
    xpath = dest.with_suffix(".xlsx")
    fast.download(spec["download_url"], xpath)
    wb = openpyxl.load_workbook(xpath, read_only=True)
    _xlsx_to_csv(wb, spec, dest, xpath.name)
    return dest


# ------------------------------------------------------------ zip of several same-header csvs
def fetch_zip_multi_csv(spec: dict, dest: Path) -> list:
    """Download a zip, extract every csv member to <sid>_members/, return the
    sorted list of paths. USAspending archives split one year into 1M-row files
    that share a header; the fast loader takes the list and checks the headers."""
    import zipfile
    mdir = dest.with_name(dest.stem + "_members")
    if mdir.exists() and any(mdir.glob("*.csv")):
        paths = sorted(mdir.glob("*.csv"))
        log(f"    members cached: {len(paths)} files")
        return paths
    zpath = dest.with_suffix(".zip")
    fast.download(spec["download_url"], zpath)
    mdir.mkdir(parents=True, exist_ok=True)
    paths = []
    with zipfile.ZipFile(zpath) as z:
        for n in sorted(z.namelist()):
            if not n.lower().endswith(".csv"):
                continue
            out = mdir / Path(n).name
            with z.open(n) as src, open(out, "wb") as dst:
                for block in iter(lambda: src.read(1 << 24), b""):
                    dst.write(block)
            paths.append(out)
    log(f"    extracted {len(paths)} members, {sum(x.stat().st_size for x in paths)/1e9:.2f} GB")
    if not spec.get("keep_zip"):
        zpath.unlink(missing_ok=True)
    return paths


# ------------------------------------------------------------ driver
def fetch(spec: dict) -> Path:
    sid = spec["source_id"]
    dest = SCRATCH / f"{sid}.csv"
    kind = spec["kind"]
    if kind == "zip_member_ranged":
        return fetch_zip_member(spec["download_url"], spec["member"], dest)
    if kind == "dol_api_paged":
        return fetch_dol_paged(spec, dest)
    if kind == "zip_xlsx":
        return fetch_zip_xlsx(spec, dest)
    if kind == "url_xlsx":
        return fetch_url_xlsx(spec, dest)
    if kind == "zip_multi_csv":
        return fetch_zip_multi_csv(spec, dest)
    if kind == "url_csv":
        fast.download(spec["download_url"], dest)   # streamed, Range resume, size-checked
        return dest
    raise RuntimeError(f"unknown kind {kind}")


def count_rows(path: Path, spec: dict) -> tuple[int, int, int, list[str], list[str]]:
    """csv-module count of the local file: rows, distinct key, null key, header, one sample row."""
    kc = spec["key_cols"][0]["col"]
    n = nulls = 0
    seen = set()
    sample = None
    header = None
    paths = list(path) if isinstance(path, (list, tuple)) else [path]
    for one in paths:
        with open(one, encoding=spec.get("encoding", "utf-8-sig"), errors="replace", newline="") as f:
            rdr = csv.reader(f, delimiter=spec.get("delimiter", ","),
                             quoting=csv.QUOTE_NONE if spec.get("quote_none") else csv.QUOTE_MINIMAL)
            h = next(rdr)
            if header is None:
                header = h
                ki = [x.strip().lower() for x in header].index(kc.strip().lower())
            for rec in rdr:
                if not rec:
                    continue
                n += 1
                v = rec[ki].strip()
                if not v:
                    nulls += 1
                else:
                    seen.add(v)
                if sample is None:
                    sample = rec
    return n, len(seen), nulls, header, sample


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sid", required=True)
    ap.add_argument("--run", action="store_true")
    ap.add_argument("--keep", action="store_true")
    args = ap.parse_args()
    fast.SCRATCH = SCRATCH
    spec = {s["source_id"]: s for s in SPECMOD.SPECS + SPECMOD2.SPECS}[args.sid]
    if spec.get("loader") != "phase5":
        raise SystemExit(f"{args.sid} is a {spec.get('loader')} spec; use that loader")
    log(f"==> {args.sid}")
    path = fetch(spec)
    if spec.get("repair_rows"):
        path = repair_rows(path, spec)
    n, d, nulls, header, sample = count_rows(path, spec)
    log(f"    local csv: {n:,} rows, {len(header)} cols, key {spec['key_cols'][0]['col']}: {d:,} distinct, {nulls:,} null")
    if not args.run:
        print("header:", header[:12], "...")
        print("sample:", (sample or [])[:12])
        print("\nadd --run to land")
        return 0
    conn = snow.connect()
    try:
        if bf._has_success(conn, args.sid):
            log("    already landed, skip")
            return 0
        r = fast.land(conn, spec, path)
        log(f"<== {args.sid} {r['status']} rows={r['rows']:,}")
        if r["status"] == "success" and not args.keep:
            for one in (path if isinstance(path, (list, tuple)) else [path]):
                Path(one).unlink(missing_ok=True)
        return 0 if r["status"] == "success" else 1
    finally:
        conn.close()


if __name__ == "__main__":
    raise SystemExit(main())
