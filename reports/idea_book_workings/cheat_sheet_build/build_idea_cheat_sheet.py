"""Stitch part_1..9.json into the idea cheat sheet: one .md in Chris's shape, one flat .xlsx."""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
REPO = Path(r"C:\Code\Ripple_v6")
HANDBOOK = REPO / "THE_WONDER_HANDBOOK.md"
OUT_MD = REPO / "reports" / "idea_cheat_sheet_2026-09-18.md"
OUT_XLSX = REPO / "reports" / "idea_cheat_sheet_2026-09-18.xlsx"

# handbook order is the order of the sheet
book_ids = re.findall(r"^### (\S+) · ", HANDBOOK.read_text(encoding="utf-8"), flags=re.M)

ideas = {}
for n in range(1, 10):
    for row in json.loads((HERE / f"part_{n}.json").read_text(encoding="utf-8")):
        ideas[row["id"]] = row

missing = [i for i in book_ids if i not in ideas]
extra = [i for i in ideas if i not in book_ids]
if missing or extra:
    print("MISSING", missing, "EXTRA", extra)
    sys.exit(1)

# HMDA codes: readers disagreed, the handbook is silent. One wording everywhere, from the public HMDA code list.
HMDA = {
    "ACTION_TAKEN": "what happened to the application: 1 loan made, 2 approved but not taken, 3 denied, 4 withdrawn, 5 file incomplete, 6 loan bought from another lender, 7 and 8 preapproval outcomes. Source: public HMDA code list, not the handbook",
    "DENIAL_REASON_1": "main reason for a denial: 1 debt too high for income, 2 job history, 3 credit history, 4 collateral, 5 not enough cash, 6 info could not be verified, 7 application incomplete, 8 mortgage insurance denied, 9 other. Source: public HMDA code list, not the handbook",
}
HMDA_PURPOSE_OLD = "what the loan was for: 1 buy a home, 2 home improvement, 3 refinance. Source: public HMDA code list, not the handbook"
HMDA_PURPOSE_NEW = "what the loan was for: 1 buy a home, 2 home improvement, 31 refinance, 32 cash-out refinance, 4 other, 5 not applicable. Source: public HMDA code list, not the handbook"
hmda_patched = 0
for it in ideas.values():
    for t in it["tables"]:
        if "HMDA" not in t["table"]:
            continue
        for c in t["columns"]:
            if c["name"] in HMDA:
                c["plain"] = HMDA[c["name"]]
                hmda_patched += 1
            elif c["name"] == "LOAN_PURPOSE":
                c["plain"] = HMDA_PURPOSE_OLD if "HISTORIC" in t["table"] else HMDA_PURPOSE_NEW
                hmda_patched += 1
print(f"HMDA code lines set to one wording: {hmda_patched}")

# two FEMA acronyms a reader left unexplained; program names are FEMA's public ones
FEMA = {
    "IHP_AMOUNT": "total dollars FEMA approved for the household under its Individuals and Households Program, the main aid pot",
    "HA_AMOUNT": "dollars FEMA approved for housing help: rent, repairs or replacement. A slice of the IHP total",
}
for it in ideas.values():
    for t in it["tables"]:
        if "FEMA" in t["table"]:
            for c in t["columns"]:
                if c["name"] in FEMA:
                    c["plain"] = FEMA[c["name"]]

# skeptic fixes 2026-09-18
for it in ideas.values():
    for t in it["tables"]:
        if not t.get("connects"):
            t["connects"] = "nothing. One table, no join"
        for c in t["columns"]:
            if c["name"] == "ASSISTANCE_TYPE_CODE" and "Source:" not in c["plain"]:
                c["plain"] = c["plain"].rstrip(". ") + ". Source: public USAspending code list, not the handbook"
            if it["id"] == "W120" and c["name"] == "SECTION_OF_ACT" and "MARTS" in t["table"]:
                c["watch"] = "0.56% filled, 17,223 of 3,087,265 rows in the 2026-09-08 profile. This is the idea's group-by key, so count it first"

# every column name must sit inside its own handbook entry, else a reader invented it
book_text = HANDBOOK.read_text(encoding="utf-8")
starts = [(m.group(1), m.start()) for m in re.finditer(r"^### (\S+) · ", book_text, flags=re.M)]
end_of_entries = book_text.find("\n# The datasets")
entry_text = {}
for i, (bid, pos) in enumerate(starts):
    stop = starts[i + 1][1] if i + 1 < len(starts) else end_of_entries
    entry_text[bid] = book_text[pos:stop]
grade = {}
for bid, txt in entry_text.items():
    m = re.search("[*][*]Data grade[.][*][*] *([ABCD])[., :;]", txt)
    grade[bid] = m.group(1) if m else "?"
print("grades", {g: list(grade.values()).count(g) for g in "ABCD?"})
bad_cols, empty_tables = [], []
for bid in book_ids:
    for t in ideas[bid]["tables"]:
        if not t["columns"]:
            empty_tables.append((bid, t["table"]))
        for c in t["columns"]:
            if c["name"] not in entry_text[bid]:
                bad_cols.append((bid, t["table"], c["name"]))
print(f"column names not found in their handbook entry: {len(bad_cols)}")
for b in bad_cols:
    print("  ", b)
print(f"tables with no columns: {len(empty_tables)}")
for e in empty_tables:
    print("  ", e)

# ---- markdown, his shape ----
lines = [
    "# Idea cheat sheet, 2026-09-18",
    "",
    "125 ideas from THE_WONDER_HANDBOOK.md, same order. Each one: the hunch, why care, the tables, the columns in plain words.",
    "Columns listed are the ones the idea needs, not every column on the table.",
    "Nothing here was run. Table and column names are copied from the handbook.",
    "Grade is the handbook's: A clean shared id, B works with a stated limit, C proxy or partial, D cannot be answered as worded.",
    "",
]
section = None
for n, bid in enumerate(book_ids, 1):
    it = ideas[bid]
    if it["section"] != section:
        section = it["section"]
        lines += [f"# {section.upper()}", ""]
    lines += [f"## {n}) {it['idea']}  `{bid}` grade {grade[bid]}", f"   - {it['why']}", ""]
    for t in it["tables"]:
        lines.append(f"**{t['table']}**")
        lines.append(f"one row = {t['one_row']}")
        if t.get("connects"):
            lines.append(f"connects: {t['connects']}")
        if not t["columns"]:
            lines.append(" - no columns named. The handbook marks this table not usable for this idea.")
        for c in t["columns"]:
            lines.append(f" - `{c['name']}`  [{c['use']}]")
            lines.append(f"      - {c['plain']}")
            if c.get("watch"):
                lines.append(f"      - WATCH: {c['watch']}")
        lines.append("")
    lines += ["---", ""]
OUT_MD.write_text("\n".join(lines), encoding="utf-8")

# ---- excel: one flat filterable sheet plus an idea index ----
wb = Workbook()
ws = wb.active
ws.title = "columns"
head = ["#", "id", "grade", "section", "idea", "why care", "table", "one row is", "connects by", "column", "plain english", "use", "watch out"]
ws.append(head)
nrows = 0
for n, bid in enumerate(book_ids, 1):
    it = ideas[bid]
    for t in it["tables"]:
        if not t["columns"]:
            ws.append([n, bid, grade[bid], it["section"], it["idea"], it["why"], t["table"], t["one_row"], t.get("connects", ""),
                       "", "no columns named; handbook marks this table not usable for this idea", "", ""])
        for c in t["columns"]:
            ws.append([n, bid, grade[bid], it["section"], it["idea"], it["why"], t["table"], t["one_row"],
                       t.get("connects", ""), c["name"], c["plain"], c["use"], c.get("watch", "")])
            nrows += 1

ix = wb.create_sheet("ideas")
ix.append(["#", "id", "grade", "section", "idea", "why care", "tables", "columns"])
for n, bid in enumerate(book_ids, 1):
    it = ideas[bid]
    ix.append([n, bid, grade[bid], it["section"], it["idea"], it["why"], len(it["tables"]),
               sum(len(t["columns"]) for t in it["tables"])])

fill = PatternFill("solid", fgColor="1F2937")
for sheet, widths in ((ws, [5, 9, 7, 9, 50, 50, 55, 36, 50, 38, 60, 9, 50]), (ix, [5, 9, 7, 9, 70, 80, 8, 9])):
    for i, w in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(OUT_XLSX)

ntables = sum(len(it["tables"]) for it in ideas.values())
print(f"ideas {len(book_ids)}  table blocks {ntables}  column rows {nrows}")
print(OUT_MD)
print(OUT_XLSX)
